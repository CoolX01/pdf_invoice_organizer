from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook


SOURCE_DIR = Path(__file__).resolve().parents[1] / "source"
sys.path.insert(0, str(SOURCE_DIR))

from exporters.excel_exporter import ExcelExporter
from models.invoice_record import InvoiceRecord
from utils.constants import EXCEL_HEADERS


def record(index: int, amount: float | None, **kwargs) -> InvoiceRecord:
    base = InvoiceRecord(
        file_path=Path(f"invoice-{index}.pdf"),
        file_name=f"invoice-{index}.pdf",
        index=index,
        invoice_number=f"1234567890123456789{index}",
        total_amount=amount,
        invoice_date="2026-05-01",
        seller_name="上海样例服务有限公司",
        buyer_name="北京示例科技有限公司",
        item_name="*信息技术服务*软件服务",
        invoice_title=f"invoice-{index}",
        parse_success=True,
    )
    for key, value in kwargs.items():
        setattr(base, key, value)
    return base


class ExcelExporterTests(unittest.TestCase):
    def test_default_export_contains_headers_rows_and_two_totals(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output = Path(tmp_dir) / "summary.xlsx"
            records = [
                record(1, 100.0),
                record(2, 200.0, duplicate_invoice=True),
                record(3, None, parse_success=False),
            ]

            ExcelExporter().export(records, output)

            workbook = load_workbook(output)
            self.assertIn("问题汇总", workbook.sheetnames)
            sheet = workbook["发票汇总"]
            headers = [sheet.cell(row=1, column=col).value for col in range(1, len(EXCEL_HEADERS) + 1)]
            self.assertEqual(headers, EXCEL_HEADERS)
            self.assertEqual(sheet.cell(row=2, column=2).value, records[0].invoice_number)
            self.assertEqual(sheet.cell(row=2, column=3).value, 100.0)
            self.assertEqual(sheet.cell(row=5, column=2).value, "金额合计（全部）")
            self.assertEqual(sheet.cell(row=5, column=3).value, 300.0)
            self.assertEqual(sheet.cell(row=6, column=2).value, "建议合计（排除重复/失败）")
            self.assertEqual(sheet.cell(row=6, column=3).value, 100.0)

            summary_sheet = workbook["问题汇总"]
            summary_values = {
                summary_sheet.cell(row=row, column=1).value: summary_sheet.cell(row=row, column=2).value
                for row in range(1, summary_sheet.max_row + 1)
            }
            self.assertEqual(summary_values["总文件数"], 3)
            self.assertEqual(summary_values["失败"], 1)
            self.assertEqual(summary_values["疑似重复发票"], 1)
            self.assertEqual(summary_values["建议合计（排除重复/失败）"], 100.0)

    def test_text_fields_are_escaped_to_prevent_excel_formula_injection(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output = Path(tmp_dir) / "summary.xlsx"
            records = [
                record(
                    1,
                    88.0,
                    invoice_number="+12345678901234567890",
                    seller_name='=HYPERLINK("https://example.invalid")',
                    buyer_name="@恶意公司",
                    item_name="-恶意项目",
                    remarks="+恶意备注",
                )
            ]

            ExcelExporter().export(records, output)

            sheet = load_workbook(output)["发票汇总"]
            self.assertEqual(sheet.cell(row=2, column=2).value, "'+12345678901234567890")
            self.assertTrue(str(sheet.cell(row=2, column=5).value).startswith("'="))
            self.assertTrue(str(sheet.cell(row=2, column=6).value).startswith("'@"))
            self.assertTrue(str(sheet.cell(row=2, column=7).value).startswith("'-"))
            self.assertTrue(str(sheet.cell(row=2, column=9).value).startswith("'+"))


if __name__ == "__main__":
    unittest.main()
