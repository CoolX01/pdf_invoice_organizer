from __future__ import annotations

from copy import copy
from datetime import datetime
import os
from pathlib import Path
import tempfile
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models.invoice_record import InvoiceRecord
from utils.constants import EXCEL_HEADERS


class ExcelExporter:
    DUPLICATE_ROW_FILL = PatternFill("solid", fgColor="FFF2CC")
    FORMULA_PREFIXES = ("=", "+", "-", "@")

    def export(
        self,
        records: list[InvoiceRecord],
        output_path: Path,
        template_path: Optional[Path] = None,
    ) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        if template_path and template_path.exists():
            workbook = load_workbook(template_path)
            sheet = workbook.active
            self._write_to_template_sheet(sheet, records)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "发票汇总"
            self._write_to_default_sheet(sheet, records)

        self._save_workbook_atomically(workbook, output_path)
        return output_path

    def _write_to_default_sheet(self, sheet, records: list[InvoiceRecord]) -> None:
        header_fill = PatternFill("solid", fgColor="DCE6F1")
        header_font = Font(bold=True)
        border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )
        alignment = Alignment(horizontal="center", vertical="center")

        for column_index, header in enumerate(EXCEL_HEADERS, start=1):
            cell = sheet.cell(row=1, column=column_index, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = alignment

        self._write_rows(sheet, records, start_row=2, border=border)
        self._set_default_column_widths(sheet)

    def _write_to_template_sheet(self, sheet, records: list[InvoiceRecord]) -> None:
        header_row = self._find_header_row(sheet)
        if header_row is None:
            self._write_to_default_sheet(sheet, records)
            return

        header_map = {}
        for column_index in range(1, sheet.max_column + 1):
            header_value = sheet.cell(row=header_row, column=column_index).value
            if isinstance(header_value, str):
                header_map[header_value.strip()] = column_index

        missing_headers = [header for header in EXCEL_HEADERS if header not in header_map]
        if missing_headers:
            self._write_to_default_sheet(sheet, records)
            return

        data_row = min(header_row + 1, max(sheet.max_row, header_row + 1))
        data_styles = {
            header: copy(sheet.cell(row=data_row, column=column_index)._style)
            for header, column_index in header_map.items()
            if sheet.max_row >= data_row
        }
        default_border = copy(sheet.cell(row=header_row, column=header_map[EXCEL_HEADERS[0]]).border)
        template_row_height = sheet.row_dimensions[data_row].height
        rows_to_delete = sheet.max_row - header_row
        if rows_to_delete > 0:
            sheet.delete_rows(header_row + 1, rows_to_delete)
        self._write_rows(
            sheet,
            records,
            start_row=header_row + 1,
            header_map=header_map,
            data_styles=data_styles,
            border=default_border,
            template_row_height=template_row_height,
        )

    def _write_rows(
        self,
        sheet,
        records: list[InvoiceRecord],
        start_row: int,
        border=None,
        header_map: Optional[dict[str, int]] = None,
        data_styles: Optional[dict[str, object]] = None,
        template_row_height: Optional[float] = None,
    ) -> None:
        header_map = header_map or {header: idx for idx, header in enumerate(EXCEL_HEADERS, start=1)}

        for row_offset, record in enumerate(records):
            row_index = start_row + row_offset
            is_duplicate_record = self._is_duplicate_record(record)
            if template_row_height is not None:
                sheet.row_dimensions[row_index].height = template_row_height
            values = self._record_to_row(record)
            for header, value in values.items():
                column_index = header_map[header]
                cell = sheet.cell(row=row_index, column=column_index, value=self._safe_cell_value(value))
                if data_styles and header in data_styles:
                    cell._style = copy(data_styles[header])
                if border:
                    cell.border = border

                if header == "价税合计（元）" and value is not None:
                    cell.number_format = "#,##0.00"
                if header == "开票日期" and value:
                    try:
                        cell.value = datetime.strptime(str(value), "%Y-%m-%d")
                        cell.number_format = "yyyy-mm-dd"
                    except ValueError:
                        cell.value = value
                if header != "备注":
                    cell.alignment = Alignment(vertical="center")
                if is_duplicate_record:
                    cell.fill = copy(self.DUPLICATE_ROW_FILL)

        total_row = start_row + len(records)
        total_label_col = header_map["发票号码"]
        total_amount_col = header_map["价税合计（元）"]
        label_cell = sheet.cell(row=total_row, column=total_label_col, value="金额合计（全部）")
        amount_cell = sheet.cell(
            row=total_row,
            column=total_amount_col,
            value=self._sum_amount(records),
        )
        if template_row_height is not None:
            sheet.row_dimensions[total_row].height = template_row_height
        label_cell.font = Font(bold=True)
        amount_cell.font = Font(bold=True)
        amount_cell.number_format = "#,##0.00"
        if border:
            label_cell.border = border
            amount_cell.border = border

        suggested_total_row = total_row + 1
        suggested_label_cell = sheet.cell(
            row=suggested_total_row,
            column=total_label_col,
            value="建议合计（排除重复/失败）",
        )
        suggested_amount_cell = sheet.cell(
            row=suggested_total_row,
            column=total_amount_col,
            value=self._sum_suggested_amount(records),
        )
        if template_row_height is not None:
            sheet.row_dimensions[suggested_total_row].height = template_row_height
        suggested_label_cell.font = Font(bold=True)
        suggested_amount_cell.font = Font(bold=True)
        suggested_amount_cell.number_format = "#,##0.00"
        if border:
            suggested_label_cell.border = border
            suggested_amount_cell.border = border

    def _find_header_row(self, sheet) -> Optional[int]:
        required_headers = set(EXCEL_HEADERS)
        max_scan_rows = min(sheet.max_row, 30)
        for row_index in range(1, max_scan_rows + 1):
            row_values = {
                str(sheet.cell(row=row_index, column=column_index).value).strip()
                for column_index in range(1, sheet.max_column + 1)
                if sheet.cell(row=row_index, column=column_index).value is not None
            }
            if required_headers.issubset(row_values):
                return row_index
        return None

    def _record_to_row(self, record: InvoiceRecord) -> dict[str, object]:
        return {
            "序号": record.index,
            "发票号码": record.invoice_number,
            "价税合计（元）": record.total_amount,
            "开票日期": record.invoice_date,
            "销售方名称": record.seller_name,
            "购买方名称": record.buyer_name,
            "项目名称": record.item_name,
            "发票名称": record.invoice_title,
            "备注": record.remarks,
        }

    def _sum_amount(self, records: list[InvoiceRecord]) -> float:
        return round(
            sum(record.total_amount or 0 for record in records),
            2,
        )

    def _sum_suggested_amount(self, records: list[InvoiceRecord]) -> float:
        return round(
            sum(
                record.total_amount or 0
                for record in records
                if record.parse_success and not self._is_duplicate_record(record)
            ),
            2,
        )

    def _is_duplicate_record(self, record: InvoiceRecord) -> bool:
        return bool(record.duplicate_invoice or record.duplicate_file)

    def _safe_cell_value(self, value):
        if not isinstance(value, str):
            return value
        if not value:
            return value
        stripped = value.lstrip()
        if stripped.startswith(self.FORMULA_PREFIXES):
            return "'" + value
        return value

    def _save_workbook_atomically(self, workbook, output_path: Path) -> None:
        temp_name = ""
        try:
            handle, temp_name = tempfile.mkstemp(
                prefix=f".{output_path.stem}.",
                suffix=output_path.suffix or ".xlsx",
                dir=output_path.parent,
            )
            os.close(handle)
            temp_path = Path(temp_name)
            workbook.save(temp_path)
            temp_path.replace(output_path)
        finally:
            if temp_name:
                temp_path = Path(temp_name)
                if temp_path.exists():
                    try:
                        temp_path.unlink()
                    except OSError:
                        pass

    def _set_default_column_widths(self, sheet) -> None:
        widths = {
            "A": 8,
            "B": 22,
            "C": 16,
            "D": 14,
            "E": 28,
            "F": 28,
            "G": 28,
            "H": 18,
            "I": 42,
        }
        for column_letter, width in widths.items():
            sheet.column_dimensions[column_letter].width = width

        for column_index in range(1, sheet.max_column + 1):
            column_letter = get_column_letter(column_index)
            sheet.column_dimensions[column_letter].bestFit = True
